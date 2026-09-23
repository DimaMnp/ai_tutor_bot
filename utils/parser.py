import io
from pypdf import PdfReader
from pptx import Presentation
from docx import Document as DocxDocument

async def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    text = ""
    ext = filename.split(".")[-1].lower()

    if ext == "pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"

    elif ext in ["pptx", "ppt"]:
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"

    elif ext in ["docx", "doc"]:
        doc = DocxDocument(io.BytesIO(file_bytes))
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text += paragraph.text + "\n"

    elif ext == "txt":
        text = file_bytes.decode("utf-8", errors="ignore")

    return text.strip()